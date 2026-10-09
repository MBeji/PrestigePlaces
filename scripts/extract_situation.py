#!/usr/bin/env python3
"""Extrait du classeur « Situation » les plans d'étage et les effectifs agrégés (sans données nominatives).

Usage : python3 -I scripts/extract_situation.py <Situation.xlsx> data/situation.json

Le classeur attendu contient :
- un onglet par niveau (RDC, 1, 2, 3, 4) où chaque position est une cellule portant un code « GROUPE.catégorie »
  (catégorie d = directeur, m = manager, c = collaborateur, e = externe) ;
- un onglet « Effetifs » (liste RH, une ligne par personne, colonne X = code de groupe, colonne V = site) ;
- un onglet « Récap » dont la colonne I porte les recrutements et externes prévus par code.
Seuls des agrégats sont écrits : aucun nom, matricule ni date individuelle ne sort du script.
"""
import collections
import json
import sys

import openpyxl

FLOORS = ['RDC', '1', '2', '3', '4']
TYPE_BY_FILL = {
    'theme:5:0.6': 'office',   # bureaux de direction (orange)
    'theme:9:0.6': 'room',     # salles et espaces communs (vert clair)
    'FF92D050': 'green',       # espaces communs (vert vif)
    'theme:7:0.8': 'legend',
    'FFFF0000': 'red',
    'FFFFFF00': 'frame',       # cadre jaune du dessin
    'theme:4:0.8': 'free',     # poste colorié sans affectation
    'theme:8:0.8': 'free',
    'theme:6:0.4': 'free',
    'FFFFC000': 'free_m',      # cellule manager sans code (extension du bureau)
}
GROUP_LABEL = {
    'BLI': 'Ammar – BLI', 'SN3': 'Ammar – SN3', 'AMMAR AUTRES': 'Ammar – autres', 'AGAL': 'Boubaker', 'Z': 'Zeineb',
    'AMINE': 'Amine', 'OMEA': 'Béji – OMEA', 'PFS': 'Béji – PFS', 'BEJI AUTRES': 'Béji – autres', 'SUP': 'Fonctions support',
}
GROUP_DIR = {
    'BLI': 'AMMAR', 'SN3': 'AMMAR', 'AMMAR AUTRES': 'AMMAR', 'AGAL': 'BOUBAKER', 'Z': 'ZEINEB', 'AMINE': 'AMINE',
    'OMEA': 'BEJI', 'PFS': 'BEJI', 'BEJI AUTRES': 'BEJI', 'SUP': 'SUPPORT',
}


def fill_key(cell):
    f = cell.fill
    if f is None or f.fill_type is None or f.fgColor is None:
        return None
    c = f.fgColor
    if c.type == 'rgb' and c.rgb and c.rgb != '00000000':
        return c.rgb
    if c.type == 'theme':
        return f'theme:{c.theme}:{round(c.tint, 2)}'
    return None


def read_floor(ws):
    cells = []
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=15, max_col=ws.max_column):
        for c in row:
            fill = fill_key(c)
            b = c.border
            has_border = any(getattr(b, side) is not None and getattr(b, side).style for side in ('left', 'right', 'top', 'bottom'))
            if c.value is not None:
                group, cat = str(c.value).rsplit('.', 1)
                cells.append({'r': c.row, 'c': c.column, 't': 'seat', 'g': group, 'k': cat})
            elif fill:
                t = TYPE_BY_FILL.get(fill, 'other')
                if t in ('legend', 'red', 'frame', 'other'):
                    continue
                cells.append({'r': c.row, 'c': c.column, 't': t})
            elif has_border:
                cells.append({'r': c.row, 'c': c.column, 't': 'wall'})
    return cells


def main(path, out):
    wb = openpyxl.load_workbook(path, data_only=True)
    floors = {}
    for name in FLOORS:
        cells = read_floor(wb[name])
        floors[name] = {'label': 'RDC' if name == 'RDC' else f'Étage {name}', 'cells': cells}
        print(name, dict(collections.Counter(x['t'] for x in cells)))

    ws = wb['Effetifs']
    headers = {ws.cell(13, c).value: c for c in range(1, 25) if ws.cell(13, c).value}
    site_col, code_col = headers['Site'], 24
    counts = collections.Counter()
    for r in range(14, ws.max_row + 1):
        if ws.cell(r, site_col).value == 'Tunis' and ws.cell(r, code_col).value:
            counts[ws.cell(r, code_col).value] += 1

    ws = wb['Récap']
    projection = {}
    for r in range(6, ws.max_row + 1):
        code, proj = ws.cell(r, 3).value, ws.cell(r, 9).value
        if code and isinstance(proj, (int, float)) and proj:
            projection[code] = proj

    groups = {}
    for g, label in GROUP_LABEL.items():
        groups[g] = {
            'label': label, 'direction': GROUP_DIR[g],
            'cdi': {k: counts.get(f'{g}.{k}', 0) for k in 'dmc'},
            'recrutements': projection.get(f'{g}.c', 0), 'externes': projection.get(f'{g}.e', 0),
        }
    data = {'source': f'{path.split("/")[-1]} — effectifs Tunis, hors stagiaires, nettoyage et sécurité', 'groups': groups, 'floors': floors}
    with open(out, 'w', encoding='utf-8') as fh:
        json.dump(data, fh, ensure_ascii=False, indent=0)
    print('écrit', out, 'positions :', sum(1 for f in floors.values() for c in f['cells'] if c['t'] == 'seat'))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
